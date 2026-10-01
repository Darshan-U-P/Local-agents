from pathlib import Path

from pptx import Presentation

from backend.planner.presentation_planner import (
    PresentationPlanner,
)

from backend.ir.builder import (
    PresentationIRBuilder,
)

from backend.renderers.pptx_renderer import (
    PPTXRenderer,
)


def main():
    print()
    print("===== PPTX RENDERER TEST =====")
    print()

    # -----------------------------------------------------
    # 1. Generate presentation plan
    # -----------------------------------------------------

    planner = PresentationPlanner()

    plan = planner.create_plan(
        topic="Quantum Computing",
        slide_count=6,
    )

    print("Planner: OK")

    # -----------------------------------------------------
    # 2. Build Presentation IR
    # -----------------------------------------------------

    builder = PresentationIRBuilder()

    presentation = builder.build(plan)

    print("Presentation IR: OK")
    print(
        f"Slides: {presentation.slide_count}"
    )
    print(
        f"Assets: {len(presentation.assets)}"
    )

    # -----------------------------------------------------
    # 3. Render PPTX
    # -----------------------------------------------------

    renderer = PPTXRenderer()

    output_path = (
        Path("generated")
        / "quantum_computing_test.pptx"
    )

    renderer.render(
        presentation,
        str(output_path),
    )

    print()
    print(
        f"PPTX generated: {output_path}"
    )

    # -----------------------------------------------------
    # 4. Verify file
    # -----------------------------------------------------

    if not output_path.exists():
        raise FileNotFoundError(
            "PPTX file was not created."
        )

    print("File exists: OK")

    # -----------------------------------------------------
    # 5. Re-open generated PPTX
    # -----------------------------------------------------

    prs = Presentation(
        str(output_path)
    )

    print(
        f"Slides in PPTX: {len(prs.slides)}"
    )

    if len(prs.slides) != presentation.slide_count:
        raise ValueError(
            "PPTX slide count does not match "
            "Presentation IR."
        )

    # -----------------------------------------------------
    # 6. Inspect slides
    # -----------------------------------------------------

    for index, slide in enumerate(
        prs.slides,
        start=1,
    ):
        print(
            f"Slide {index}: "
            f"{len(slide.shapes)} shapes"
        )

    print()
    print("===== PPTX VALIDATION =====")
    print("PPTX GENERATED AND VALID")
    print()


if __name__ == "__main__":
    main()