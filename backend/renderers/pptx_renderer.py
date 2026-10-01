from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from backend.ir.presentation_ir import PresentationIR
from backend.layout.geometry import (
    SLIDE_WIDTH,
    SLIDE_HEIGHT,
    PositionedElement,
)


class PPTXRenderer:
    """
    Converts Presentation IR + Layout Engine output
    into an editable PowerPoint presentation.
    """

    def __init__(self):
        pass

    # ---------------------------------------------------------
    # PUBLIC API
    # ---------------------------------------------------------

    def render(
        self,
        presentation: PresentationIR,
        output_path: str,
    ) -> str:
        """
        Render a PresentationIR into an editable .pptx file.
        """

        presentation.validate()

        output = Path(output_path)
        output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        prs = Presentation()

        # 16:9 widescreen
        prs.slide_width = Inches(SLIDE_WIDTH)
        prs.slide_height = Inches(SLIDE_HEIGHT)

        # Remove the default empty slide if one exists.
        if len(prs.slides) > 0:
            slide_id = prs.slides._sldIdLst[0]
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)

        for slide_ir in presentation.slides:
            self._render_slide(
                prs,
                presentation,
                slide_ir,
            )

        prs.save(str(output))

        return str(output)

    # ---------------------------------------------------------
    # SLIDE
    # ---------------------------------------------------------

    def _render_slide(
        self,
        prs,
        presentation,
        slide_ir,
    ):
        """
        Create one PowerPoint slide.
        """

        blank_layout = prs.slide_layouts[6]

        slide = prs.slides.add_slide(blank_layout)

        # Background
        self._set_background(
            slide,
            presentation.theme.background_color,
        )

        # Layout engine
        from backend.layout.layout_engine import LayoutEngine

        layout_engine = LayoutEngine()

        positioned_elements = layout_engine.layout_slide(
            slide_ir
        )

        # Render every positioned element.
        for element in positioned_elements:
            self._render_element(
                slide,
                element,
                presentation,
            )

    # ---------------------------------------------------------
    # ELEMENT ROUTER
    # ---------------------------------------------------------

    def _render_element(
        self,
        slide,
        element: PositionedElement,
        presentation,
    ):
        """
        Route positioned elements to the correct renderer.
        """

        if element.element_type == "text":
            self._render_text(
                slide,
                element,
                presentation,
            )
            return

        if element.role == "visual":
            self._render_visual(
                slide,
                element,
                presentation,
            )
            return

        # Unknown element type.
        self._render_placeholder(
            slide,
            element,
            "Unsupported element",
        )

    # ---------------------------------------------------------
    # TEXT
    # ---------------------------------------------------------

    def _render_text(
        self,
        slide,
        element,
        presentation,
    ):
        """
        Render an editable PowerPoint text box.
        """

        rect = element.rect

        textbox = slide.shapes.add_textbox(
            Inches(rect.x),
            Inches(rect.y),
            Inches(rect.width),
            Inches(rect.height),
        )

        text_frame = textbox.text_frame

        text_frame.clear()

        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        paragraph = text_frame.paragraphs[0]

        paragraph.text = str(element.content)

        # Alignment
        alignment = "left"

        if element.style:
            alignment = element.style.alignment

        if alignment == "center":
            paragraph.alignment = PP_ALIGN.CENTER

        elif alignment == "right":
            paragraph.alignment = PP_ALIGN.RIGHT

        else:
            paragraph.alignment = PP_ALIGN.LEFT

        # Font
        run = paragraph.runs[0]

        if element.style:
            run.font.size = Pt(
                element.style.font_size
            )

            run.font.bold = (
                element.style.bold
            )

        else:
            run.font.size = Pt(20)

        run.font.name = (
            presentation.theme.font_family
        )

        # Text color
        run.font.color.rgb = self._hex_to_rgb(
            presentation.theme.text_color
        )

    # ---------------------------------------------------------
    # VISUAL
    # ---------------------------------------------------------

    def _render_visual(
        self,
        slide,
        element,
        presentation,
    ):
        """
        Render an asset.

        If a real asset exists:
            render the actual image.

        If the asset does not exist yet:
            render an editable placeholder.

        This allows the PPTX pipeline to work before
        the image-generation system is implemented.
        """

        asset = self._find_asset(
            presentation,
            element.asset_id,
        )

        if asset and asset.path:
            asset_path = Path(asset.path)

            if asset_path.exists():
                self._render_image(
                    slide,
                    element,
                    asset_path,
                )
                return

        description = element.content

        if not description:
            description = "Visual asset"

        self._render_placeholder(
            slide,
            element,
            str(description),
        )

    # ---------------------------------------------------------
    # IMAGE
    # ---------------------------------------------------------

    def _render_image(
        self,
        slide,
        element,
        image_path: Path,
    ):
        """
        Add a real image to the slide.
        """

        rect = element.rect

        slide.shapes.add_picture(
            str(image_path),
            Inches(rect.x),
            Inches(rect.y),
            width=Inches(rect.width),
            height=Inches(rect.height),
        )

    # ---------------------------------------------------------
    # PLACEHOLDER
    # ---------------------------------------------------------

    def _render_placeholder(
        self,
        slide,
        element,
        text,
    ):
        """
        Create an editable placeholder shape.

        Used until actual image/diagram/chart
        generation is implemented.
        """

        rect = element.rect

        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(rect.x),
            Inches(rect.y),
            Inches(rect.width),
            Inches(rect.height),
        )

        # Fill
        shape.fill.solid()

        shape.fill.fore_color.rgb = RGBColor(
            235,
            238,
            245,
        )

        # Border
        shape.line.color.rgb = RGBColor(
            150,
            160,
            180,
        )

        # Text
        text_frame = shape.text_frame

        text_frame.clear()

        text_frame.word_wrap = True
        text_frame.vertical_anchor = (
            MSO_ANCHOR.MIDDLE
        )

        paragraph = text_frame.paragraphs[0]

        paragraph.text = (
            f"Visual placeholder\n\n{text}"
        )

        paragraph.alignment = PP_ALIGN.CENTER

        for run in paragraph.runs:
            run.font.size = Pt(14)
            run.font.name = "Aptos"
            run.font.color.rgb = RGBColor(
                80,
                90,
                110,
            )

    # ---------------------------------------------------------
    # BACKGROUND
    # ---------------------------------------------------------

    def _set_background(
        self,
        slide,
        color,
    ):
        """
        Set slide background color.
        """

        fill = slide.background.fill

        fill.solid()

        fill.fore_color.rgb = self._hex_to_rgb(
            color
        )

    # ---------------------------------------------------------
    # ASSET LOOKUP
    # ---------------------------------------------------------

    @staticmethod
    def _find_asset(
        presentation,
        asset_id,
    ):
        if asset_id is None:
            return None

        for asset in presentation.assets:
            if asset.id == asset_id:
                return asset

        return None

    # ---------------------------------------------------------
    # COLOR
    # ---------------------------------------------------------

    @staticmethod
    def _hex_to_rgb(hex_color):
        """
        Convert:
            #2563EB

        into:
            RGBColor(37, 99, 235)
        """

        value = hex_color.strip().lstrip("#")

        if len(value) != 6:
            raise ValueError(
                f"Invalid hex color: {hex_color}"
            )

        red = int(value[0:2], 16)
        green = int(value[2:4], 16)
        blue = int(value[4:6], 16)

        return RGBColor(
            red,
            green,
            blue,
        )